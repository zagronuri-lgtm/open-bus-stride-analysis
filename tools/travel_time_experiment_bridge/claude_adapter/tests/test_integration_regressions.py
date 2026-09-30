"""Synthetic local regressions for exclusive output ownership and consistent offsets."""
import os
import tempfile
import unittest
from unittest.mock import patch
import openpyxl
from test_prepare import base_trips, make_template, job_for, manifest
import prepare_update_trips as adapter

class IntegrationRegressions(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.template = os.path.join(self.tmp.name, 'template.xlsx')
        self.out = os.path.join(self.tmp.name, 'out')
        self.trips = base_trips()
        self.sha = make_template(self.template, self.trips)
        self.job = job_for(self.trips, {'10100_1_0_06:00': 428}, self.sha)

    def test_interleaved_same_digest_has_one_writer(self):
        original = adapter._revalidate
        losers = []
        def interleave(job):
            # Winner owns a directory but has not created any artefact yet.
            loser = adapter.prepare(job, self.template, manifest(self.sha), self.out, drop_columns=False)
            losers.append(loser)
            self.assertEqual(loser.status, 'blocked')
            self.assertEqual(os.listdir(loser.job_dir), [])
            return original(job)
        with patch.object(adapter, '_revalidate', side_effect=interleave):
            winner = adapter.prepare(self.job, self.template, manifest(self.sha), self.out)
        self.assertEqual(len(losers), 1)
        self.assertEqual(winner.status, 'prepared', winner.blocks)
        self.assertEqual(winner.output_sha256, adapter.sha256_file(winner.output_path))
        self.assertIsNone(losers[0].report_path)

    def test_preexisting_empty_directory_is_not_owned_or_cleaned(self):
        dest = os.path.join(self.out, self.job.digest)
        os.makedirs(dest)
        result = adapter.prepare(self.job, self.template, manifest(self.sha), self.out)
        self.assertEqual(result.status, 'blocked')
        self.assertTrue(os.path.isdir(dest))
        self.assertEqual(os.listdir(dest), [])

    def test_missing_or_blank_offset_blocks_before_write(self):
        for index, value in enumerate((None, '', '   ')):
            with self.subTest(value=value):
                wb = openpyxl.load_workbook(self.template)
                ws = wb['Trips']; headers = [c.value for c in ws[1]]
                ws.cell(2, headers.index('Day Offset') + 1).value = value
                wb.save(self.template); wb.close()
                sha = adapter.sha256_file(self.template)
                job = job_for(self.trips, {'10100_1_0_06:00': 428}, sha)
                result = adapter.prepare(job, self.template, manifest(sha), self.out + str(index))
                self.assertEqual(result.status, 'blocked')
                self.assertTrue(any('Day Offset' in b for b in result.blocks))
                self.assertFalse(os.path.exists(os.path.join(result.job_dir, adapter.OUTPUT_NAME)))
                self.assertTrue(os.path.isfile(result.report_path))

    def test_postwrite_exception_removes_only_owned_output_and_reports_block(self):
        sentinel = os.path.join(self.tmp.name, 'keep.txt')
        with open(sentinel, 'w') as fh: fh.write('untouched')
        with patch.object(adapter, 'roundtrip_check', side_effect=ValueError('injected failure')):
            result = adapter.prepare(self.job, self.template, manifest(self.sha), self.out)
        self.assertEqual(result.status, 'blocked')
        self.assertTrue(any('output write/verification failed' in b for b in result.blocks))
        self.assertFalse(os.path.exists(os.path.join(result.job_dir, adapter.OUTPUT_NAME)))
        self.assertTrue(os.path.isfile(result.report_path))
        self.assertEqual(adapter.sha256_file(self.template), self.sha)
        with open(sentinel) as fh: self.assertEqual(fh.read(), 'untouched')

if __name__ == '__main__': unittest.main()
