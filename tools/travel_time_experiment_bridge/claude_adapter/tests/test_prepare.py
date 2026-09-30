"""Synthetic tests for claude-adapter/prepare_update_trips.py. No network, no Optibus. The template built here mirrors
the real EXPORT TRIPS workbook observed on 30.9.2026 (sheet names, columns, time cells as datetime.time, Day Offset text)."""
import datetime as dt
import hashlib
import json
import os
import sys
import tempfile
import unittest

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, '..'))
sys.path.insert(0, os.path.join(HERE, '..', '..', 'implementation'))
from prepare_update_trips import prepare, sha256_file, verify_diff, AdapterError, cell_clock, CellError, OUTPUT_NAME  # noqa: E402
from bridge_contract import ValidatedJob  # noqa: E402
from bridge_contract import validate_job, ContractError  # noqa: E402

TRIPS_HEADERS = ['Id', 'Region', 'Catalog Number', 'Sign', 'Direction', 'Alternative', 'Origin Stop id', 'Destination Stop Id', 'Day Offset',
                 'Departure', 'Arrival', 'Vehicle Type Ids', 'Distance', 'Existing', 'Custom', 'Days', 'Boarding Time', 'Offboarding Time',
                 'Sub trip index', 'Route Id', 'Vehicle ID', 'customer']
STOP_HEADERS = ['Trip Id', 'Time', 'Point Id', 'distance', 'Sequence', 'Sequence Id', 'Time Point', 'tp']


def t(m):
    return dt.time((m % 1440) // 60, m % 60)


def make_template(path, trips, stops=None, days='1'):
    """trips: list of dict(id, makat, direction, alt, dep, arr, dist, origin, dest); stops: optional override per id (list of (min, point, leg))."""
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    ws = wb.create_sheet('Trips'); ws.append(TRIPS_HEADERS)
    st = wb.create_sheet('Places'); st.append(['Id', 'Description', 'Address', 'Latitude', 'Longitude', 'Type'])
    wb.create_sheet('Stops').append(['Id', 'Description', 'Place', 'Address', 'Latitude', 'Longitude', 'Type'])
    ss = wb.create_sheet('StopTimes'); ss.append(STOP_HEADERS)
    wb.create_sheet('ReliefPoints').append(['Sign', 'Stop id'])
    vt = wb.create_sheet('VehicleTypes'); vt.append(['Id', 'Description', 'Family Type']); vt.append(['interurban', 'interurban', 'None'])
    wb.create_sheet('Taxis').append(['Origin Stop Id', 'Destination Stop Id'])
    pr = wb.create_sheet('Parameters'); pr.append(['Name', 'Value']); pr.append(['allow_duplicate_base_route_key', False])
    mp = wb.create_sheet('TripIdsMapping'); mp.append(['System Id', 'User Id'])
    places = {}
    for tr in trips:
        off = '1' if tr['dep'] >= 1440 else '0'
        ws.append([tr['id'], None, tr['makat'], tr['makat'][-3:], tr['direction'], tr['alt'], tr['origin'], tr['dest'], off, t(tr['dep']), t(tr['arr']),
                   'interurban', tr['dist'], None, None, days, 0, 0, None, f"{tr['makat']}-{tr['direction']}-{tr['alt']}", None, 'metropoline'])
        rows = (stops or {}).get(tr['id']) or [(tr['dep'], tr['origin'], 0), (tr['arr'], tr['dest'], tr['dist'])]
        for i, (m, point, leg) in enumerate(rows):
            ss.append([tr['id'], t(m), point, leg, i, None, True, True])
            places[point] = 1
        mp.append([tr['id'], tr['id']])
    for p in sorted(places):
        st.append([p, f'stop {p}', None, 32.0, 34.8, 'depot' if p.startswith('99') else 'stop'])
    wb.save(path)
    return sha256_file(path)


def base_trips():
    return [dict(id='10100_1_0_06:00', makat='10100', direction='1', alt='0', dep=360, arr=420, dist=20.5, origin='99991', dest='11111'),
            dict(id='10100_2_0_23:30', makat='10100', direction='2', alt='0', dep=1410, arr=1460, dist=20.5, origin='11111', dest='99991'),
            dict(id='10277_1_0_07:00', makat='10277', direction='1', alt='0', dep=420, arr=470, dist=30.0, origin='99991', dest='22222'),
            dict(id='10100_1_0_25:00', makat='10100', direction='1', alt='0', dep=1500, arr=1540, dist=20.5, origin='99991', dest='11111')]


def job_for(trips, changes, sha, sched='copy-1', ds='ds-copy', rev='rev-1', selection=None):
    src = dict(dataset_id=ds, schedule_id=sched, revision=rev, sha256=sha)
    common = dict(schema_version=1, source=src, branch='synthetic', day_type='weekday', season='autumn', service_dates=['2026-09-30'],
                  engine_version='engine-1', policy_version='policy-1')
    snapshot = dict(common, trips=[])
    payload = dict(common, run_id='run-1', selection_trip_ids=selection if selection is not None else sorted(changes), trips=[])
    for tr in trips:
        line = tr['makat'][2:]
        ident = dict(id=tr['id'], operator='op', makat=tr['makat'], line_number=line, direction=tr['direction'], alternative=tr['alt'],
                     departure=tr['dep'], before_arrival=tr['arr'])
        after = changes.get(tr['id'], tr['arr'])
        snapshot['trips'].append(dict(ident, evidence=dict(status='change_allowed' if tr['id'] in changes else 'review', recommended_arrival=after,
                                                           verified_service_dates=['2026-09-30'], verified_day_type='weekday', reason='synthetic')))
        payload['trips'].append(dict(ident, after_arrival=after))
    return validate_job(payload, snapshot)


def manifest(sha, sched='copy-1', ds='ds-copy', rev='rev-1', origin='origin-1', days='1'):
    return dict(dataset_id=ds, schedule_id=sched, revision=rev, sha256=sha, origin_schedule_id=origin, days=days,
                exported_at='2026-09-30T07:29:00', export_method='UI EXPORT TRIPS')


class PrepareTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.tpl = os.path.join(self.tmp, 'template.xlsx')
        self.trips = base_trips()
        self.sha = make_template(self.tpl, self.trips)
        self.out = os.path.join(self.tmp, 'out')

    def test_happy_path_changes_only_arrival_and_last_time(self):
        job = job_for(self.trips, {'10100_1_0_06:00': 428}, self.sha)
        r = prepare(job, self.tpl, manifest(self.sha), self.out)
        self.assertEqual(r.status, 'prepared', r.blocks)
        self.assertEqual(r.trips_total, 4); self.assertEqual(r.trips_changed, 1); self.assertEqual(r.minutes_delta_sum, 8)
        self.assertEqual(r.diff_summary, {'Trips.Arrival': 1, 'Trips.__dropped__': ['Custom', 'customer'], 'StopTimes.Time': 1, 'StopTimes.__dropped__': ['tp']})
        wb = openpyxl.load_workbook(r.output_path)
        ws = wb['Trips']; h = [c.value for c in ws[1]]
        self.assertNotIn('customer', h); self.assertNotIn('Custom', h)
        row = next(rw for rw in ws.iter_rows(min_row=2, values_only=True) if rw[0] == '10100_1_0_06:00')
        self.assertEqual(row[h.index('Arrival')], dt.time(7, 8)); self.assertEqual(row[h.index('Departure')], dt.time(6, 0))
        self.assertTrue(os.path.exists(r.report_path)); self.assertEqual(r.output_sha256, sha256_file(r.output_path))
        self.assertEqual(os.path.basename(r.output_path), OUTPUT_NAME); self.assertEqual(os.path.basename(r.job_dir), job.digest)
        self.assertEqual(r.roundtrip['mismatches'], []); self.assertEqual(r.roundtrip['checked'], 4)
        # template untouched
        self.assertEqual(sha256_file(self.tpl), self.sha)

    def test_unchanged_job_prepares_identical_data(self):
        job = job_for(self.trips, {}, self.sha, selection=['10100_1_0_06:00'])
        r = prepare(job, self.tpl, manifest(self.sha), self.out, drop_columns=False)
        self.assertEqual(r.status, 'prepared', r.blocks); self.assertEqual(r.diff_summary, {})

    def test_midnight_wrap_after_change(self):
        job = job_for(self.trips, {'10100_2_0_23:30': 1475}, self.sha)   # 23:30 -> 00:35 next day, offset stays 0
        r = prepare(job, self.tpl, manifest(self.sha), self.out)
        self.assertEqual(r.status, 'prepared', r.blocks)
        ws = openpyxl.load_workbook(r.output_path)['Trips']; h = [c.value for c in ws[1]]
        row = next(rw for rw in ws.iter_rows(min_row=2, values_only=True) if rw[0] == '10100_2_0_23:30')
        self.assertEqual(row[h.index('Arrival')], dt.time(0, 35)); self.assertEqual(row[h.index('Day Offset')], '0')

    def test_after_midnight_departure_offset_one(self):
        job = job_for(self.trips, {'10100_1_0_25:00': 1545}, self.sha)
        r = prepare(job, self.tpl, manifest(self.sha), self.out)
        self.assertEqual(r.status, 'prepared', r.blocks)
        ws = openpyxl.load_workbook(r.output_path)['Trips']; h = [c.value for c in ws[1]]
        row = next(rw for rw in ws.iter_rows(min_row=2, values_only=True) if rw[0] == '10100_1_0_25:00')
        self.assertEqual(row[h.index('Arrival')], dt.time(1, 45)); self.assertEqual(row[h.index('Day Offset')], '1')

    def test_block_source_equals_copy(self):
        job = job_for(self.trips, {'10100_1_0_06:00': 428}, self.sha)
        r = prepare(job, self.tpl, manifest(self.sha, origin='copy-1'), self.out)
        self.assertEqual(r.status, 'blocked'); self.assertTrue(any('source == copy' in b for b in r.blocks)); self.assertIsNone(r.output_path)
        self.assertFalse(os.path.exists(os.path.join(r.job_dir, OUTPUT_NAME)))

    def test_block_protected_map(self):
        job = job_for(self.trips, {'10100_1_0_06:00': 428}, self.sha)
        r = prepare(job, self.tpl, manifest(self.sha), self.out, protected_schedule_ids=('copy-1',))
        self.assertEqual(r.status, 'blocked'); self.assertTrue(any('protected' in b for b in r.blocks))

    def test_block_hash_mismatch(self):
        job = job_for(self.trips, {'10100_1_0_06:00': 428}, 'b' * 64)
        r = prepare(job, self.tpl, manifest('b' * 64), self.out)
        self.assertEqual(r.status, 'blocked'); self.assertTrue(any('sha256' in b for b in r.blocks))

    def test_block_revision_mismatch(self):
        job = job_for(self.trips, {'10100_1_0_06:00': 428}, self.sha, rev='rev-2')
        r = prepare(job, self.tpl, manifest(self.sha), self.out)
        self.assertEqual(r.status, 'blocked'); self.assertTrue(any('revision' in b for b in r.blocks))

    def test_block_departure_mismatch(self):
        trips = base_trips(); trips[0]['dep'] = 365   # payload thinks departure differs from the template
        job = job_for(trips, {'10100_1_0_06:00': 428}, self.sha)
        r = prepare(job, self.tpl, manifest(self.sha), self.out)
        self.assertEqual(r.status, 'blocked'); self.assertTrue(any('departure' in b for b in r.blocks))

    def test_block_missing_and_extra_trips(self):
        job = job_for(self.trips[:3], {'10100_1_0_06:00': 428}, self.sha)
        r = prepare(job, self.tpl, manifest(self.sha), self.out)
        self.assertEqual(r.status, 'blocked'); self.assertTrue(any('not in payload' in b for b in r.blocks))

    def test_block_intermediate_stops_on_changed_trip(self):
        tpl2 = os.path.join(self.tmp, 't2.xlsx')
        sha2 = make_template(tpl2, self.trips, stops={'10100_1_0_06:00': [(360, '99991', 0), (390, '55555', 10.0), (420, '11111', 10.5)]})
        job = job_for(self.trips, {'10100_1_0_06:00': 428}, sha2)
        r = prepare(job, tpl2, manifest(sha2), self.out)
        self.assertEqual(r.status, 'blocked'); self.assertTrue(any('intermediate' in b for b in r.blocks))
        job2 = job_for(self.trips, {'10277_1_0_07:00': 0} if False else {'10100_1_0_25:00': 1545}, sha2)   # unchanged 3-stop trip is tolerated
        r2 = prepare(job2, tpl2, manifest(sha2), self.out)
        self.assertEqual(r2.status, 'prepared', r2.blocks)

    def test_block_distance_mismatch(self):
        tpl3 = os.path.join(self.tmp, 't3.xlsx')
        sha3 = make_template(tpl3, self.trips, stops={'10100_1_0_06:00': [(360, '99991', 0), (420, '11111', 25.0)]})
        job = job_for(self.trips, {'10100_1_0_06:00': 428}, sha3)
        r = prepare(job, tpl3, manifest(sha3), self.out)
        self.assertEqual(r.status, 'blocked'); self.assertTrue(any('distance' in b for b in r.blocks))

    def test_block_days_mismatch(self):
        job = job_for(self.trips, {'10100_1_0_06:00': 428}, self.sha)
        r = prepare(job, self.tpl, manifest(self.sha, days='6'), self.out)
        self.assertEqual(r.status, 'blocked'); self.assertTrue(any('Days' in b for b in r.blocks))

    def test_contract_rejects_decrease_before_adapter(self):
        with self.assertRaises(ContractError):
            job_for(self.trips, {'10100_1_0_06:00': 400}, self.sha)

    def test_contract_freezes_277_1(self):
        with self.assertRaises(ContractError):
            job_for(self.trips, {'10277_1_0_07:00': 480}, self.sha)

    def test_bad_manifest_raises(self):
        job = job_for(self.trips, {'10100_1_0_06:00': 428}, self.sha)
        with self.assertRaises(AdapterError):
            prepare(job, self.tpl, {'schedule_id': 'x'}, self.out)

    def test_verify_diff_detects_foreign_edit(self):
        job = job_for(self.trips, {'10100_1_0_06:00': 428}, self.sha)
        r = prepare(job, self.tpl, manifest(self.sha), self.out)
        wb = openpyxl.load_workbook(r.output_path); wb['Trips'].cell(3, 10).value = dt.time(23, 59); wb.save(r.output_path)
        d = verify_diff(self.tpl, r.output_path)
        self.assertEqual(d.get('Trips.Departure'), 1)




class ReviewRound1Regressions(unittest.TestCase):
    """One regression per review finding (claude-adapter-review-round1.md)."""
    def setUp(self):
        self.tmp = tempfile.mkdtemp(); self.tpl = os.path.join(self.tmp, 'template.xlsx'); self.trips = base_trips()
        self.sha = make_template(self.tpl, self.trips); self.out = os.path.join(self.tmp, 'out')

    def _job(self, changes=None, **kw):
        return job_for(self.trips, changes if changes is not None else {'10100_1_0_06:00': 428}, self.sha, **kw)

    # 1. run_id never becomes a path component; existing outputs are never overwritten
    def test_run_id_path_traversal_is_content_only(self):
        trips = self.trips; sha = self.sha
        src = dict(dataset_id='ds-copy', schedule_id='copy-1', revision='rev-1', sha256=sha)
        common = dict(schema_version=1, source=src, branch='b', day_type='weekday', season='autumn', service_dates=['2026-09-30'], engine_version='e', policy_version='p')
        snap = dict(common, trips=[]); pay = dict(common, run_id='../../escape', selection_trip_ids=['10100_1_0_06:00'], trips=[])
        for tr in trips:
            ident = dict(id=tr['id'], operator='op', makat=tr['makat'], line_number=tr['makat'][2:], direction=tr['direction'], alternative=tr['alt'], departure=tr['dep'], before_arrival=tr['arr'])
            after = 428 if tr['id'] == '10100_1_0_06:00' else tr['arr']
            snap['trips'].append(dict(ident, evidence=dict(status='change_allowed' if after != tr['arr'] else 'review', recommended_arrival=after, verified_service_dates=['2026-09-30'], verified_day_type='weekday', reason='x')))
            pay['trips'].append(dict(ident, after_arrival=after))
        job = validate_job(pay, snap)
        r = prepare(job, self.tpl, manifest(sha), self.out)
        self.assertEqual(r.status, 'prepared', r.blocks)
        self.assertTrue(os.path.abspath(r.output_path).startswith(os.path.abspath(self.out) + os.sep))
        self.assertNotIn('escape', r.output_path)
        with open(r.report_path, encoding='utf-8') as fh:
            self.assertEqual(json.load(fh)['run_id'], '../../escape')

    def test_existing_job_dir_is_not_overwritten(self):
        job = self._job(); r1 = prepare(job, self.tpl, manifest(self.sha), self.out)
        self.assertEqual(r1.status, 'prepared'); before = sha256_file(r1.output_path)
        r2 = prepare(job, self.tpl, manifest(self.sha), self.out)
        self.assertEqual(r2.status, 'blocked'); self.assertTrue(any('not overwriting' in b for b in r2.blocks)); self.assertEqual(sha256_file(r1.output_path), before)

    # 2. times: no modulo, strict decoding, round-trip, duration >= 24h blocked
    def test_block_duration_over_24h(self):
        r = prepare(self._job({'10100_1_0_06:00': 360 + 1440}), self.tpl, manifest(self.sha), self.out)
        self.assertEqual(r.status, 'blocked'); self.assertTrue(any('>= 1440' in b for b in r.blocks))

    def test_block_seconds_and_bad_clock_text(self):
        with self.assertRaises(CellError): cell_clock(dt.time(6, 0, 30))
        with self.assertRaises(CellError): cell_clock('25:00')
        with self.assertRaises(CellError): cell_clock('06:60')
        with self.assertRaises(CellError): cell_clock('06:00:30')
        with self.assertRaises(CellError): cell_clock(0.25 + 1e-4)
        with self.assertRaises(CellError): cell_clock(True)
        self.assertEqual(cell_clock('06:05'), 365); self.assertEqual(cell_clock(0.25), 360); self.assertEqual(cell_clock(dt.time(23, 59)), 1439)

    def test_template_time_with_seconds_blocks(self):
        wb = openpyxl.load_workbook(self.tpl); ws = wb['Trips']; h = [c.value for c in ws[1]]
        ws.cell(2, h.index('Departure') + 1).value = dt.time(6, 0, 30); wb.save(self.tpl); sha = sha256_file(self.tpl)
        r = prepare(job_for(self.trips, {'10100_1_0_06:00': 428}, sha), self.tpl, manifest(sha), self.out)
        self.assertEqual(r.status, 'blocked'); self.assertTrue(any('seconds' in b for b in r.blocks))

    def test_roundtrip_reconstructs_after_midnight_and_offset(self):
        r = prepare(self._job({'10100_2_0_23:30': 1475, '10100_1_0_25:00': 1550}), self.tpl, manifest(self.sha), self.out)
        self.assertEqual(r.status, 'prepared', r.blocks); self.assertEqual(r.roundtrip['mismatches'], [])

    # 3. Sign identity
    def test_block_sign_mismatch(self):
        wb = openpyxl.load_workbook(self.tpl); ws = wb['Trips']; h = [c.value for c in ws[1]]
        ws.cell(2, h.index('Sign') + 1).value = '999'; wb.save(self.tpl); sha = sha256_file(self.tpl)
        r = prepare(job_for(self.trips, {'10100_1_0_06:00': 428}, sha), self.tpl, manifest(sha), self.out)
        self.assertEqual(r.status, 'blocked'); self.assertTrue(any('Sign' in b for b in r.blocks))

    # 4. distances and sequences
    def test_block_missing_distance_even_on_unchanged_trip(self):
        wb = openpyxl.load_workbook(self.tpl); ws = wb['Trips']; h = [c.value for c in ws[1]]
        ws.cell(4, h.index('Distance') + 1).value = None; wb.save(self.tpl); sha = sha256_file(self.tpl)
        r = prepare(job_for(self.trips, {'10100_1_0_06:00': 428}, sha), self.tpl, manifest(sha), self.out)
        self.assertEqual(r.status, 'blocked'); self.assertTrue(any('Distance missing' in b for b in r.blocks))

    def test_block_nan_leg_distance(self):
        wb = openpyxl.load_workbook(self.tpl); ws = wb['StopTimes']; h = [c.value for c in ws[1]]
        ws.cell(3, h.index('distance') + 1).value = float('nan'); wb.save(self.tpl); sha = sha256_file(self.tpl)
        r = prepare(job_for(self.trips, {'10100_1_0_06:00': 428}, sha), self.tpl, manifest(sha), self.out)
        self.assertEqual(r.status, 'blocked'); self.assertTrue(any('leg missing or not a finite' in b for b in r.blocks))

    def test_block_sequence_gap_and_duplicate(self):
        wb = openpyxl.load_workbook(self.tpl); ws = wb['StopTimes']; h = [c.value for c in ws[1]]
        ws.cell(3, h.index('Sequence') + 1).value = 5; wb.save(self.tpl); sha = sha256_file(self.tpl)
        r = prepare(job_for(self.trips, {'10100_1_0_06:00': 428}, sha), self.tpl, manifest(sha), self.out)
        self.assertEqual(r.status, 'blocked'); self.assertTrue(any('Sequence must be 0..n-1' in b for b in r.blocks))
        wb = openpyxl.load_workbook(self.tpl); ws = wb['StopTimes']; ws.cell(3, h.index('Sequence') + 1).value = 0; wb.save(self.tpl); sha = sha256_file(self.tpl)
        r = prepare(job_for(self.trips, {'10100_1_0_06:00': 428}, sha), self.tpl, manifest(sha), self.out)
        self.assertEqual(r.status, 'blocked'); self.assertTrue(any('Sequence must be 0..n-1' in b for b in r.blocks))

    def test_block_orphan_stoptimes_row(self):
        wb = openpyxl.load_workbook(self.tpl); ws = wb['StopTimes']; ws.append(['ghost_trip', dt.time(1, 0), '1', 0, 0, None, True, True]); wb.save(self.tpl); sha = sha256_file(self.tpl)
        r = prepare(job_for(self.trips, {'10100_1_0_06:00': 428}, sha), self.tpl, manifest(sha), self.out)
        self.assertEqual(r.status, 'blocked'); self.assertTrue(any('orphan' in b for b in r.blocks))

    # 5. duplicate headers
    def test_block_duplicate_header(self):
        wb = openpyxl.load_workbook(self.tpl); ws = wb['Trips']; ws.cell(1, ws.max_column + 1).value = 'Arrival'; wb.save(self.tpl); sha = sha256_file(self.tpl)
        r = prepare(job_for(self.trips, {'10100_1_0_06:00': 428}, sha), self.tpl, manifest(sha), self.out)
        self.assertEqual(r.status, 'blocked'); self.assertTrue(any('duplicate header' in b for b in r.blocks))

    # 6. trust boundary: a hand-built ValidatedJob is re-validated and rejected
    def test_block_forged_validated_job(self):
        real = self._job()
        forged = ValidatedJob(real.run_id, 'f' * 64, real.payload_json, real.snapshot_json)
        r = prepare(forged, self.tpl, manifest(self.sha), self.out)
        self.assertEqual(r.status, 'blocked'); self.assertTrue(any('re-validation' in b for b in r.blocks))
        tampered = json.loads(real.payload_json); tampered['trips'][0]['after_arrival'] = 999
        forged2 = ValidatedJob(real.run_id, real.digest, json.dumps(tampered), real.snapshot_json)
        r2 = prepare(forged2, self.tpl, manifest(self.sha), self.out + '-forged2')
        self.assertEqual(r2.status, 'blocked'); self.assertTrue(any('re-validation' in b for b in r2.blocks))
        with self.assertRaises(AdapterError):
            prepare(ValidatedJob(real.run_id, 'not-hex', real.payload_json, real.snapshot_json), self.tpl, manifest(self.sha), self.out)


if __name__ == '__main__':
    unittest.main()
