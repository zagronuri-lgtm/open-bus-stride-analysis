"""Real XLSX regressions for council findings; no network or Optibus."""
import os
import tempfile
import unittest
from unittest.mock import patch
import openpyxl
from test_prepare import base_trips, make_template, job_for, manifest
import prepare_update_trips as adapter

class CouncilRegressions(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(); self.addCleanup(self.tmp.cleanup)
        self.template = os.path.join(self.tmp.name, 'template.xlsx')
        self.out = os.path.join(self.tmp.name, 'out')
        self.trips = base_trips(); self.sha = make_template(self.template, self.trips)

    def job(self):
        return job_for(self.trips, {'10100_1_0_06:00': 428}, self.sha)

    def mutate(self, sheet, column, row, value):
        wb = openpyxl.load_workbook(self.template); ws = wb[sheet]
        headers = [c.value for c in ws[1]]; ws.cell(row, headers.index(column)+1).value = value
        wb.save(self.template); wb.close(); self.sha = adapter.sha256_file(self.template)

    def test_template_replacement_cannot_change_frozen_approved_cells(self):
        job = self.job(); original_load = openpyxl.load_workbook; replaced = []
        def load(source, *args, **kwargs):
            if not replaced:
                replaced.append(True)
                wb = original_load(self.template); ws = wb['Trips']; headers = [c.value for c in ws[1]]
                ws.cell(2, headers.index('Vehicle Type Ids')+1).value = 'unapproved-replacement'
                wb.save(self.template); wb.close()
            return original_load(source, *args, **kwargs)
        with patch.object(adapter.openpyxl, 'load_workbook', side_effect=load):
            result = adapter.prepare(job, self.template, manifest(self.sha), self.out)
        self.assertEqual(result.status, 'prepared', result.blocks)
        self.assertEqual(result.template_sha256, self.sha)
        self.assertNotEqual(adapter.sha256_file(self.template), self.sha)
        wb = original_load(result.output_path); ws = wb['Trips']; headers = [c.value for c in ws[1]]
        self.assertEqual(ws.cell(2, headers.index('Vehicle Type Ids')+1).value, 'interurban'); wb.close()

    def test_endpoints_must_match_even_when_times_match(self):
        for row in (2, 3):
            with self.subTest(row=row):
                make_template(self.template, self.trips)
                self.mutate('StopTimes', 'Point Id', row, 'wrong-stop')
                result = adapter.prepare(self.job(), self.template, manifest(self.sha), self.out+str(row))
                self.assertEqual(result.status, 'blocked')
                self.assertTrue(any('Point Id' in b for b in result.blocks))
                self.assertIsNone(result.output_path)

    def test_endpoint_headers_required(self):
        wb = openpyxl.load_workbook(self.template); ws = wb['Trips']; headers = [c.value for c in ws[1]]
        ws.cell(1, headers.index('Origin Stop id')+1).value = 'Wrong origin header'
        wb.save(self.template); wb.close(); self.sha = adapter.sha256_file(self.template)
        result = adapter.prepare(self.job(), self.template, manifest(self.sha), self.out)
        self.assertEqual(result.status, 'blocked'); self.assertTrue(any('Origin Stop id' in b for b in result.blocks))

    def test_bad_zip_returns_block_report_without_workbook(self):
        with open(self.template, 'wb') as stream: stream.write(b'not an xlsx archive')
        self.sha = adapter.sha256_file(self.template)
        result = adapter.prepare(self.job(), self.template, manifest(self.sha), self.out)
        self.assertEqual(result.status, 'blocked'); self.assertTrue(any('BadZipFile' in b for b in result.blocks))
        self.assertTrue(os.path.isfile(result.report_path))
        self.assertFalse(os.path.exists(os.path.join(result.job_dir, adapter.OUTPUT_NAME)))

    def test_negative_leg_blocked_even_with_matching_sum(self):
        self.mutate('StopTimes', 'distance', 2, -10)
        self.mutate('StopTimes', 'distance', 3, 30.5)
        result = adapter.prepare(self.job(), self.template, manifest(self.sha), self.out)
        self.assertEqual(result.status, 'blocked'); self.assertTrue(any('leg is negative' in b for b in result.blocks))

    def test_negative_trip_distance_blocked(self):
        self.mutate('Trips', 'Distance', 2, -20.5)
        self.mutate('StopTimes', 'distance', 3, -20.5)
        result = adapter.prepare(self.job(), self.template, manifest(self.sha), self.out)
        self.assertEqual(result.status, 'blocked'); self.assertTrue(any('Trips.Distance is negative' in b for b in result.blocks))

if __name__ == '__main__': unittest.main()
